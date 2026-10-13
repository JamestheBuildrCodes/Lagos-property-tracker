import csv
import tempfile
import unittest
from pathlib import Path

import client_doc_report
from report_builder import build_xlsx_from_data
from market_analysis import build_summary

ROOT = Path(__file__).resolve().parents[1]


class ClientDocReportTests(unittest.TestCase):
    def sample_rows(self):
        return [
            {"date_scraped":"2026-09-29", "market_node":"Lekki Phase 1", "location":"Lekki Phase 1",
             "source":"PropertyPro.ng", "source_url":"https://propertypro.ng/property/lekki-phase-1-example-listing",
             "transaction":"sale", "property_type":"flat_apartment", "bedrooms":"3",
             "asking_price_ngn":"700000000", "price_per_sqm_ngn":"1400000", "title":"3 bedroom apartment"},
            {"date_scraped":"2026-09-28", "market_node":"Lekki Phase 1", "location":"Lekki Phase 1",
             "source":"PropertyPro.ng", "source_url":"https://propertypro.ng/agent/bad-record",
             "transaction":"sale", "property_type":"flat_apartment", "bedrooms":"3",
             "asking_price_ngn":"700000000", "price_per_sqm_ngn":"1400000", "title":"agent profile"},
        ]

    def test_breakdown_requires_auditable_listing_url(self):
        rows = client_doc_report.bedroom_breakdown(self.sample_rows(), "Lekki Phase 1", "sale")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["sample"], "1 listing")
        self.assertEqual(rows[0]["link"], "https://propertypro.ng/property/lekki-phase-1-example-listing")

    def test_docx_contains_clickable_source_hyperlink_and_style(self):
        from market_analysis import clean_rows
        rows = clean_rows(self.sample_rows())
        summary = build_summary(rows, [], None)
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "client.docx"
            client_doc_report.build_docx(rows, summary, str(out))
            self.assertTrue(out.exists())
            from zipfile import ZipFile
            with ZipFile(out) as z:
                document_xml = z.read("word/document.xml").decode("utf-8")
                rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
            self.assertIn("Segment", document_xml)
            self.assertIn("Reported price", document_xml)
            self.assertIn("Range", document_xml)
            self.assertIn("Listings (n)", document_xml)
            self.assertIn("PropertyPro.ng", document_xml)
            self.assertIn("propertypro.ng/property/lekki-phase-1-example-listing", rels)
            self.assertIn("D9D9D9", document_xml)

    def test_xlsx_headline_currency_and_percent_formats(self):
        from market_analysis import clean_rows
        rows = clean_rows(self.sample_rows())
        summary = build_summary(rows, [], None)
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "report.xlsx"
            build_xlsx_from_data(rows, [], [], summary, str(out))
            from openpyxl import load_workbook
            wb = load_workbook(out, data_only=False)
            ws = wb["Executive Summary"]
            sale_row = next(r for r in range(1, ws.max_row+1) if ws.cell(r,1).value == "Median apartment sale")
            self.assertEqual(ws.cell(sale_row,2).value, "₦700,000,000")
            score = wb["Area Scorecard"]
            self.assertEqual(score["C2"].number_format, "#,##0")
            self.assertIn("%", score["E2"].number_format)
            self.assertNotEqual(score["E2"].number_format, "#,##0.00")


    def test_xlsx_has_explicit_estate_intel_public_research_status(self):
        from market_analysis import clean_rows
        rows = clean_rows(self.sample_rows())
        summary = build_summary(rows, [], None)
        research = [{
            "market_node": "Lekki Phase 1",
            "title": "Public research overview",
            "url": "https://estateintel.com/insights/example",
            "date_scraped": "2026-09-29",
            "public_sale_price_ngn": "",
            "land_area_sqm": "",
            "size_units": "",
        }]
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "report.xlsx"
            build_xlsx_from_data(rows, [], research, summary, str(out))
            from openpyxl import load_workbook
            wb = load_workbook(out, data_only=True)
            self.assertIn("Estate Intel Public", wb.sheetnames)
            self.assertIn("Source Health", wb.sheetnames)
            self.assertIn("Additional Sources", wb.sheetnames)
            additional = wb["Additional Sources"]
            self.assertIn("Knight Frank Nigeria", [additional.cell(i, 1).value for i in range(2, additional.max_row + 1)])
            self.assertIn("Cushman & Wakefield (C&W)", [additional.cell(i, 1).value for i in range(2, additional.max_row + 1)])
            self.assertIn("Estate Intel commercial arrangement status", [additional.cell(i, 1).value for i in range(2, additional.max_row + 1)])
            health = wb["Source Health"]
            row = next(i for i in range(2, health.max_row + 1) if health.cell(i, 1).value == "Estate Intel")
            self.assertEqual(health.cell(row, 2).value, 1)
            self.assertEqual(health.cell(row, 3).value, 0)
            self.assertIn("Public research only", health.cell(row, 6).value)

    def test_docs_writer_formats_currency_and_uses_real_bullets(self):
        import docs_writer
        from market_analysis import clean_rows
        summary = build_summary(clean_rows(self.sample_rows()), [], None)
        text = docs_writer.report_text(summary)
        self.assertIn("₦700,000,000", text)
        requests = docs_writer._format_requests(text)
        self.assertTrue(any("createParagraphBullets" in request for request in requests))
        # The API requests remove the literal glyph before applying a real bullet.
        self.assertTrue(any("deleteContentRange" in request for request in requests))


    def test_source_breakdown_matches_sample_table_shape(self):
        rows = client_doc_report._source_breakdown_rows(
            self.sample_rows(),
            "Lekki Phase 1",
            "sale",
            "3",
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["label"], "3-bedroom")
        self.assertEqual(rows[0]["value"], "₦700,000,000 (single observed asking price)")
        self.assertEqual(rows[0]["range"], "Not meaningful with 1 listing")
        self.assertEqual(rows[0]["sample"], "1")
        self.assertEqual(rows[0]["source"], "PropertyPro.ng")
        self.assertEqual(rows[0]["as_of"], "2026-09-29")


    def test_google_doc_is_created_new_and_has_requested_table_styles(self):
        source = Path(client_doc_report.__file__).read_text(encoding="utf-8")
        import re

        # Normalize whitespace so the test checks implementation intent,
        # not source formatting choices such as line wrapping.
        normalized = re.sub(r"\s+", "", source)

        self.assertIn("drive.files().create(", normalized)
        self.assertIn("GOOGLE_OAUTH_TOKEN_JSON", source)
        self.assertNotIn("GOOGLE_DOC_ID", source)
        self.assertIn('"updateTableCellStyle"', source)
        self.assertIn('"link":{"url":rows[ri-1]["link"]}', normalized)
        mailjet = Path(ROOT / "mailjet_sender.py").read_text(encoding="utf-8")
        self.assertIn("mailjet_sender", source)
        self.assertIn("send_message(", source)
        self.assertIn("MAILJET_API_KEY", mailjet)
        self.assertIn("MAILJET_SECRET_KEY", mailjet)
        self.assertIn('"1. What This Report Is"', source)
        self.assertIn('"7. What Happens Next"', source)
        self.assertIn('"8. Sources & Definitions"', source)
        self.assertIn('"Estate Intel — Public Research Context"', source)
        self.assertIn('"Range"', source)
        self.assertIn('"Listings (n)"', source)

    def test_narrative_discloses_missing_port_harcourt_and_estate_intel_context(self):
        from market_analysis import clean_rows
        rows = clean_rows(self.sample_rows())
        summary = build_summary(rows, [], None)
        narrative = client_doc_report._narrative_lines(rows, summary, research=[
            {"market_node": "Lekki Phase 1", "title": "Public research page",
             "url": "https://estateintel.com/insights/public-example",
             "date_scraped": "2026-09-29", "public_sale_price_ngn": ""}])
        rendered = "\n".join(value for kind, value in narrative if kind != "table")
        self.assertIn("Port Harcourt is not currently tracked", rendered)
        self.assertIn("Estate Intel", rendered)
        self.assertIn("commercial arrangement", rendered.lower())
        self.assertIn("08:00 Lagos time", rendered)
        self.assertIn("1. What This Report Is", rendered)
        self.assertIn("7. What Happens Next", rendered)
        self.assertIn("8. Sources & Definitions", rendered)

    def test_scraper_rejects_promo_copy_from_listing_cards(self):
        import scraper
        html = """<div class='results'>
        <div class='promo'>Premium Plus Top promoted listing. Email you when a new listing matches. Get alerts.</div>
        <div class='property-card'><a href='/property/3-bedroom-flat-for-sale-old-ikoyi-ikoyi-lagos-AAA11'>3 Bedroom Flat Old Ikoyi</a><span>₦500,000,000</span><span>Updated 14 Sep 2026</span><span>3 Beds</span></div>
        <div class='property-card'><a href='/property/4-bedroom-flat-for-sale-old-ikoyi-ikoyi-lagos-BBB22'>4 Bedroom Flat Old Ikoyi</a><span>₦700,000,000</span><span>Updated 14 Sep 2026</span><span>4 Beds</span></div>
        </div>"""
        cards = scraper.candidate_cards(html, "PropertyPro.ng", "https://propertypro.ng/property-for-sale/flat-apartment/in/lagos/ikoyi/old-ikoyi")
        self.assertEqual(len(cards), 2)
        self.assertTrue(all("premium plus" not in text.lower() for _, text in cards))
        self.assertTrue(all("email you when a new listing matches" not in text.lower() for _, text in cards))


    def test_fresh_service_account_doc_is_reused_without_oauth_duplicate(self):
        import json
        import os

        current_dir = Path.cwd()
        try:
            with tempfile.TemporaryDirectory() as td:
                os.chdir(td)
                url = "https://docs.google.com/document/d/fresh-weekly-doc/edit"
                Path("google_doc_url.txt").write_text(url + chr(10), encoding="utf-8")
                Path("google_doc_publication.json").write_text(json.dumps({
                    "publication_status": "new_weekly_document_created",
                    "created_new_document": True,
                    "url": url,
                }), encoding="utf-8")
                result = client_doc_report._use_fresh_service_account_doc()
                self.assertEqual(result, url)
                self.assertEqual(Path("client_doc_url.txt").read_text(encoding="utf-8").strip(), url)
                verification = json.loads(Path("google_doc_verification.json").read_text(encoding="utf-8"))
                self.assertEqual(verification["publication_status"], "fresh_service_account_weekly_doc_used")
                self.assertTrue(verification["document_is_current_run"])
        finally:
            os.chdir(current_dir)

    def test_revoked_oauth_token_falls_back_to_fresh_service_account_doc(self):
        import json
        import os
        from google.auth.exceptions import RefreshError

        error = RefreshError(
            "invalid_grant: Token has been expired or revoked."
        )
        self.assertTrue(client_doc_report._is_revoked_oauth_refresh_error(error))
        self.assertFalse(
            client_doc_report._is_revoked_oauth_refresh_error(
                RuntimeError("Google Docs network timeout")
            )
        )

        current_dir = Path.cwd()
        try:
            with tempfile.TemporaryDirectory() as td:
                os.chdir(td)
                fallback_url = "https://docs.google.com/document/d/layer3-report-id/edit"
                Path("google_doc_url.txt").write_text(fallback_url + chr(10), encoding="utf-8")
                Path("google_doc_publication.json").write_text(json.dumps({
                    "publication_status": "new_weekly_document_created",
                    "created_new_document": True,
                    "url": fallback_url,
                }), encoding="utf-8")

                result = client_doc_report._fallback_to_fresh_service_account_doc(error)

                self.assertEqual(result, fallback_url)
                self.assertEqual(Path("client_doc_url.txt").read_text(encoding="utf-8").strip(), fallback_url)
                verification = json.loads(
                    Path("google_doc_verification.json").read_text(encoding="utf-8")
                )
                self.assertEqual(
                    verification["publication_status"],
                    "fallback_to_fresh_service_account_weekly_doc",
                )
                self.assertFalse(verification["client_preferred_doc_created"])
                self.assertFalse(verification["client_preferred_doc_content_verified"])
                self.assertIn("GOOGLE_OAUTH_TOKEN_JSON", verification["action_required"])
        finally:
            os.chdir(current_dir)

    def test_revoked_oauth_token_without_fallback_doc_fails_with_clear_action(self):
        import os
        from google.auth.exceptions import RefreshError

        current_dir = Path.cwd()
        error = RefreshError("invalid_grant: Token has been expired or revoked.")
        try:
            with tempfile.TemporaryDirectory() as td:
                os.chdir(td)
                with self.assertRaisesRegex(RuntimeError, "no verified service-account weekly document"):
                    client_doc_report._fallback_to_fresh_service_account_doc(error)
        finally:
            os.chdir(current_dir)


    def test_revoked_oauth_without_fresh_doc_records_docx_attachment_fallback(self):
        import json
        import os
        from google.auth.exceptions import RefreshError

        current_dir = Path.cwd()
        error = RefreshError("invalid_grant: Token has been expired or revoked.")
        try:
            with tempfile.TemporaryDirectory() as td:
                os.chdir(td)
                docx = Path("Lagos_Property_Market_Snapshot_2026-10-13.docx")
                docx.write_bytes(b"test docx payload")
                marker = client_doc_report._write_docx_attachment_fallback(error, docx)
                self.assertEqual(marker, "ATTACHED_DOCX_FALLBACK")
                self.assertEqual(
                    Path("client_doc_url.txt").read_text(encoding="utf-8").strip(),
                    "ATTACHED_DOCX_FALLBACK",
                )
                verification = json.loads(
                    Path("google_doc_verification.json").read_text(encoding="utf-8")
                )
                self.assertEqual(
                    verification["publication_status"],
                    "docx_attachment_fallback_no_google_doc",
                )
                self.assertFalse(verification["google_doc_created"])
                self.assertEqual(verification["docx_attachment"], docx.name)
        finally:
            os.chdir(current_dir)



if __name__ == "__main__":
    unittest.main()
