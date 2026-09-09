"""Mirror the reporting workbook into the existing Google Sheet and style it."""
from __future__ import annotations
import json, os, sys
from pathlib import Path
import gspread
from google.oauth2.service_account import Credentials

SCOPES=["https://www.googleapis.com/auth/spreadsheets","https://www.googleapis.com/auth/drive"]
DESIRED=["Executive Summary","Area Scorecard","Current Listings","Weekly Changes","Estate Intel Public","Methodology"]

def get_client():
    raw=os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON","").strip()
    if not raw: raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON is missing")
    data=json.loads(raw); email=data.get("client_email") or "unknown"; project=data.get("project_id") or "unknown"
    print(f"Google service account: {email} | project: {project}")
    return gspread.authorize(Credentials.from_service_account_info(data,scopes=SCOPES)),email

def share_viewer(sheet_id,email):
    if not email or "@" not in email:return
    from googleapiclient.discovery import build
    creds=get_client()[0].auth
    drive=build("drive","v3",credentials=creds,cache_discovery=False)
    drive.permissions().create(fileId=sheet_id,body={"type":"user","role":"reader","emailAddress":email},sendNotificationEmail=False).execute()
    print(f"Google Sheet viewer access granted: {email}")

def run():
    if len(sys.argv)<2: raise SystemExit("Usage: python sheets_writer.py <listings.csv> [changes.csv] [research.csv] [xlsx_path]")
    listings=sys.argv[1]; changes=sys.argv[2] if len(sys.argv)>2 else "none"; research=sys.argv[3] if len(sys.argv)>3 else "none"; xlsx_path=sys.argv[4] if len(sys.argv)>4 else ""
    sheet_id=os.environ.get("GOOGLE_SHEET_ID","").strip()
    if not sheet_id: raise RuntimeError("GOOGLE_SHEET_ID is missing")
    client,service_email=get_client()
    try: sheet=client.open_by_key(sheet_id)
    except Exception as exc: raise RuntimeError(f"Could not open Google Sheet. Confirm GOOGLE_SHEET_ID and share the existing sheet with {service_email} as Editor. Original error: {exc}") from exc
    if not xlsx_path or not Path(xlsx_path).exists(): raise RuntimeError("Layer 2 workbook is missing.")
    from openpyxl import load_workbook
    wb=load_workbook(xlsx_path,data_only=False)
    for name in DESIRED:
        values=[[cell.value for cell in row] for row in wb[name].iter_rows()]
        try: ws=sheet.worksheet(name); ws.clear()
        except gspread.WorksheetNotFound: ws=sheet.add_worksheet(title=name,rows=max(len(values)+10,100),cols=max(max(len(r) for r in values),10))
        rows=max(len(values)+5,20); cols=max(max((len(r) for r in values),default=1)+2,10)
        ws.resize(rows=rows,cols=cols)
        if values: ws.update(values,"A1")
        ws.freeze(rows=3 if name=="Executive Summary" else 1)
        ws.format("A1:Z1",{"backgroundColor":{"red":0.05,"green":0.14,"blue":0.25},"textFormat":{"foregroundColor":{"red":1,"green":1,"blue":1},"bold":True},"horizontalAlignment":"LEFT","verticalAlignment":"MIDDLE"})
        if name=="Executive Summary":
            ws.format("A3:C3",{"backgroundColor":{"red":0.91,"green":0.94,"blue":0.97},"textFormat":{"bold":True,"foregroundColor":{"red":0.05,"green":0.14,"blue":0.25}}})
            ws.format("A14:C14",{"textFormat":{"bold":True,"foregroundColor":{"red":0.05,"green":0.14,"blue":0.25},"fontSize":12}})
        ws.format(f"A1:{gspread.utils.rowcol_to_a1(rows,cols)}",{"wrapStrategy":"WRAP","verticalAlignment":"TOP"})
        # Practical column widths without trying to build a dashboard.
        try:
            ws.columns_auto_resize(0, cols)
        except Exception as exc:
            print(f"Column auto-resize skipped for {name}: {exc}")
    client_email=os.environ.get("REPORT_CLIENT_EMAIL","").strip()
    if client_email:
        try: share_viewer(sheet_id,client_email)
        except Exception as exc: print(f"Client Sheet sharing skipped: {exc}")
    print(f"Google Sheets updated and styled: {sheet.title}")

if __name__=="__main__": run()
