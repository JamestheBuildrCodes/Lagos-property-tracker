"""Build the three client-facing intelligence layers from one normalized snapshot.

Layer 1: concise executive HTML email.
Layer 2: polished XLSX mirrored into the existing Google Sheet.
Layer 3: concise analyst DOCX; docs_writer.py can publish the same content to Google Docs.
"""
from __future__ import annotations
import json, os, sys
from datetime import datetime, timezone
from pathlib import Path
from market_analysis import build_summary, clean_rows, fmt_naira, load_csv

TODAY=datetime.now(timezone.utc).strftime("%Y-%m-%d")

SUPPLEMENTAL_SOURCE_ASSESSMENTS = [
    {
        "Source": "Knight Frank Nigeria",
        "Official page": "https://www.knightfrank.com.ng/research/lagos-market-update-13035.aspx",
        "Source type": "Institutional market research",
        "Assessment": "Official Lagos Market Update H1 2026 page identified.",
        "Use in report": "Contextual triangulation for market commentary and broad trends only; do not mix publication-level estimates with current listing-level medians.",
        "Next step / limitation": "Review the report period, sector definitions and methodology before comparing figures. The download flow may request contact details.",
    },
    {
        "Source": "Cushman & Wakefield (C&W)",
        "Official page": "https://www.cushmanwakefield.com/en/insights/global-data-center-market-comparison",
        "Source type": "Commercial real estate research / data centres",
        "Assessment": "Current official public research references Lagos in a global data-centre comparison; this is a commercial infrastructure segment, not a residential listing feed.",
        "Use in report": "Context only for future commercial-property research; do not blend data-centre indicators with residential sale/rental medians.",
        "Next step / limitation": "Confirm local/regional Nigeria research coverage and reuse/access permissions before including comparable commercial property figures.",
    },
    {
        "Source": "Estate Intel commercial arrangement status",
        "Official page": "https://estateintel.com/about",
        "Source type": "Subscription / commercial access status",
        "Assessment": "Current client agreement, paid subscription, or pending negotiation is not verifiable from the repository or public pages reviewed.",
        "Use in report": "The tracker currently uses public research pages only; premium/login-gated content is not collected.",
        "Next step / limitation": "Confirm status with the account owner and Estate Intel directly before claiming a commercial arrangement is active or pending.",
    },
]

def esc(v):
    return (str(v if v is not None else "").replace("&","&amp;").replace("<","&lt;")
            .replace(">","&gt;").replace('"',"&quot;"))

def _load_inputs(listings_path,changes_path="none",research_path=None):
    listings=clean_rows(load_csv(listings_path))
    changes=load_csv(changes_path) if changes_path and changes_path.lower()!="none" and os.path.exists(changes_path) else []
    research=load_csv(research_path) if research_path and research_path.lower()!="none" and os.path.exists(research_path) else []
    summary=build_summary(listings,changes,listings_path)
    try:
        quality = json.loads(Path("market_output_quality.json").read_text(encoding="utf-8"))
        summary["data_quality"] = {
            "input_rows": quality.get("input_rows"),
            "evidence_verified_rows": quality.get("evidence_verified_rows", quality.get("output_rows")),
            "quarantined_for_review": quality.get("quarantined_for_review", 0),
            "duplicates_removed": quality.get("duplicates_removed", 0),
        }
    except (OSError, json.JSONDecodeError):
        summary["data_quality"] = {}
    research_sources={"Estate Intel"} if research else set()
    if research_sources:
        summary["sources"]=len(set(summary.get("source_names",[])) | research_sources)
        summary["source_names"]=sorted(set(summary.get("source_names",[])) | research_sources)
    try:
        from ai_market_interpreter import interpret_summary
        ai_commentary=interpret_summary(summary)
        if ai_commentary:
            summary["ai_commentary"]=ai_commentary
    except Exception as exc:
        if os.environ.get("OPENAI_AI_COMMENTARY_REQUIRED", "false").strip().lower() == "true":
            raise RuntimeError(f"Required OpenAI analysis failed; refusing to publish report: {type(exc).__name__}: {exc}") from exc
        print(f"[AI analysis unavailable; using deterministic narrative] {type(exc).__name__}: {exc}")
    return listings,changes,research,summary

def _wow_text(metric):
    p=metric.get("pct_change")
    if p is None:return "No prior snapshot"
    if p>0:return f"↑ {p:.1f}%"
    if p<0:return f"↓ {abs(p):.1f}%"
    return "→ 0.0%"

def build_executive_html_from_summary(s,date=None,doc_url="",sheet_url=""):
    """Render an executive brief that explains implications and actions, not just metrics."""
    date=date or datetime.now(timezone.utc).strftime("%d %B %Y")
    wow=s.get("week_on_week",{}) or {}
    changes=s.get("changes",{}) or {}
    sig=s.get("signals",{}) or {}
    def money(v, compact=True):
        if v in (None,""): return "Not enough comparable data"
        try:
            n=float(str(v).replace(",",""))
        except (TypeError,ValueError):
            return str(v)
        if compact and abs(n)>=1_000_000_000:
            return f"₦{n/1_000_000_000:,.2f} billion"
        if compact and abs(n)>=1_000_000:
            return f"₦{n/1_000_000:,.1f} million"
        return f"₦{n:,.0f}"
    def pct(key):
        if not wow.get("available"): return "Baseline unavailable"
        metric=wow.get("metrics",{}).get(key,{})
        value=metric.get("pct_change")
        if value is None: return "Not comparable"
        direction="increased" if value>0 else "decreased" if value<0 else "was unchanged"
        return f"{direction} {abs(value):.1f}% in the tracked sample"
    def node(obj): return esc(obj.get("market_node","Area not available")) if obj else "Not enough comparable data"
    def link(url,label):
        if not url: return ""
        return f"<a href=\"{esc(url)}\" style=\"display:inline-block;padding:10px 14px;background:#0d2340;color:#ffffff;text-decoration:none;margin:4px 8px 4px 0\">{esc(label)}</a>"
    notes=s.get("analysis_notes",[]) or []
    ai_commentary=s.get("ai_commentary") or {}
    valid_count=int(s.get("listings_tracked") or 0)
    reductions=int(changes.get("price_reductions") or 0)
    increases=int(changes.get("price_increases") or 0)
    new_count=int(changes.get("new_listings") or 0)
    not_observed=int(changes.get("not_observed", changes.get("delisted", 0)) or 0)
    populated=[x for x in s.get("scorecard",[]) if int(x.get("listing_count") or 0)>0]
    thin=[x.get("market_node","Unknown") for x in sorted(populated,key=lambda x:int(x.get("listing_count") or 0))[:3]]
    # Build cautious statements from the actual data, never a generic claim about all of Lagos.
    takeaway=[]
    if valid_count==0:
        takeaway.append("No validated comparable listings passed the current quality gate. This run must not be used to draw market conclusions.")
    else:
        takeaway.append(f"This scan retained {valid_count:,} listings that passed the current source and location checks across {s.get('areas_monitored',0)} monitored locations.")
    if wow.get("available"):
        takeaway.append(f"The tracked sample shows apartment sale asking prices {pct('sale')}; annual rent asking prices {pct('rent')}; land asking price per square metre {pct('land_ppsqm')}. These movements are sample-level signals, not market-wide changes.")
    if reductions:
        takeaway.append(f"{reductions:,} listings have a recorded asking-price reduction and should be reviewed as potential negotiation leads.")
    else:
        takeaway.append("No price reductions were verified in the comparison file. Do not infer that sellers are not negotiating; the scan may not have captured a price change.")
    if increases:
        takeaway.append(f"{increases:,} listings have a recorded asking-price increase in the available comparison.")
    if not_observed:
        takeaway.append(f"{not_observed:,} previous URLs were not observed in this scan. This does not prove that the properties sold or were delisted; it may also reflect coverage or extraction differences.")
    if thin:
        takeaway.append("Sample caution: the thinnest covered locations are "+", ".join(thin)+". Treat these as watchlist observations, not reliable area-wide benchmarks.")
    dq = s.get("data_quality", {}) or {}
    if dq:
        takeaway.append(
            f"Listing-card cross-check: {dq.get('evidence_verified_rows', 0)} rows passed deterministic source-evidence checks; "
            f"{dq.get('quarantined_for_review', 0)} rows were excluded from comparable statistics for review. "
            "This check tests consistency with captured public listing text; it does not prove ownership or a closed transaction."
        )
    investment=sig.get("investment")
    buyer=sig.get("buyer_value")
    land=sig.get("land_opportunity")
    html_parts=[
      "<!doctype html><html><body style=\"font-family:Arial,sans-serif;color:#172033;line-height:1.55;max-width:760px;margin:0 auto\">",
      f"<header style=\"padding:24px;background:#0d2340;color:#fff\"><div style=\"font-size:11px;letter-spacing:1.5px\">MASTER BUILDER INC. | MARKET INTELLIGENCE</div><h1 style=\"font-size:26px;line-height:1.2;margin:10px 0 4px\">Lagos &amp; Abuja Property Market</h1><div style=\"font-size:17px\">Weekly Intelligence Brief</div><div style=\"margin-top:8px\">{esc(date)}</div></header>",
      "<main style=\"padding:18px 4px\">",
      "<h2 style=\"font-size:19px\">Executive takeaway</h2>",
      f"<p>{esc(ai_commentary.get('executive_takeaway') or 'This brief separates verified observations from tentative market signals. Use it to decide which listings and locations deserve follow-up, then use the linked property records to verify any material figure before committing funds.')}</p>",
      "<h2 style=\"font-size:19px\">Market snapshot</h2>",
      "<table role=\"presentation\" style=\"width:100%;border-collapse:collapse\"><tr style=\"background:#e8eef5\"><th align=\"left\" style=\"padding:9px\">Indicator</th><th align=\"left\" style=\"padding:9px\">This week</th><th align=\"left\" style=\"padding:9px\">Change vs previous scan</th></tr>"
    ]
    metrics=[
      ("Validated comparable listings",f"{valid_count:,}",pct("listings")),
      ("Locations monitored",f"{int(s.get('areas_monitored') or 0):,}","Coverage scope"),
      ("Sources represented",f"{int(s.get('sources') or 0):,}","See methodology"),
      ("Newly observed listings",f"{new_count:,}","Newly observed, not necessarily newly posted"),
      ("Price reductions recorded",f"{reductions:,}","Potential negotiation leads"),
      ("Price increases recorded",f"{increases:,}","Check source history"),
      ("Median apartment asking sale price",money(s.get("median_apartment_sale")),"Asking prices, not transaction prices"),
      ("Median annual apartment asking rent",money(s.get("median_annual_rent")),"Per year, where source states annual rent"),
      ("Median land asking price / sqm",money(s.get("median_land_ppsqm")),"Only valid when price-per-sqm and land area are parsed consistently"),
    ]
    for label,value,context in metrics:
        html_parts.append(f"<tr><td style=\"padding:9px;border-bottom:1px solid #d9e1ea\">{esc(label)}</td><td style=\"padding:9px;border-bottom:1px solid #d9e1ea\"><b>{esc(value)}</b></td><td style=\"padding:9px;border-bottom:1px solid #d9e1ea\">{esc(context)}</td></tr>")
    html_parts.append("</table>")
    html_parts.append("<p><b>Reporting and distribution:</b> The automated scan is scheduled weekly for Mondays at 08:00 Lagos time (07:00 UTC). Client email is dispatched only after quality checks and report publication succeed. The workflow checks the email provider receipt; that confirms provider acceptance, not guaranteed inbox placement.</p>")
    html_parts.append("<h2 style=\"font-size:19px\">What changed this week?</h2>")
    if ai_commentary.get("key_findings"):
        html_parts.append("<ul>")
        html_parts.extend(f"<li style=\"margin:6px 0\">{esc(x)}</li>" for x in ai_commentary["key_findings"])
        html_parts.append("</ul>")
    else:
        html_parts.append("<ul>"+ "".join(f"<li style=\"margin:6px 0\">{esc(n)}</li>" for n in notes[:5]) +"</ul>" if notes else "<p>No market narrative was produced because the available evidence is insufficient.</p>")
    html_parts.append("<h2 style=\"font-size:19px\">What this means for the business</h2><ul>")
    html_parts.extend(f"<li style=\"margin:7px 0\">{esc(x)}</li>" for x in takeaway)
    if investment:
        y=investment.get("gross_rent_yield_proxy_pct")
        if y is not None:
            html_parts.append(f"<li><b>Rental screening:</b> {node(investment)} has an indicative gross asking-rent yield proxy of {float(y):.2f}%. It is a screening ratio, not net return, and should be relied on only with a sufficiently large and comparable sample.</li>")
    if buyer:
        sample=int(buyer.get("sale_count") or 0)
        html_parts.append(f"<li><b>Buyer screening:</b> {node(buyer)} is the current lowest observed sale-price-per-square-metre screen among qualifying locations (sale sample: {sample}). Check that property type, condition, size and tenure are comparable before interpreting this as value.</li>")
    if land:
        html_parts.append(f"<li><b>Land screening:</b> {node(land)} has the lowest observed land asking price per sqm in the current sample. This is a lead for due diligence, not a recommendation to buy.</li>")
    if ai_commentary.get("business_implications"):
        html_parts.append("<h2 style=\"font-size:19px\">Business implications</h2><ul>")
        html_parts.extend(f"<li style=\"margin:7px 0\">{esc(x)}</li>" for x in ai_commentary["business_implications"])
        html_parts.append("</ul>")
    html_parts.append("</ul><h2 style=\"font-size:19px\">Recommended actions</h2><ol>")
    actions=[]
    if reductions: actions.append(f"Review the {reductions:,} recorded price reductions first and confirm the seller's current price directly.")
    if not_observed: actions.append("Treat previously observed URLs missing this week as 'not observed'; verify status directly before calling any property sold or delisted.")
    actions.append("Compare like-for-like listings by location, transaction type, bedroom count, size and source update date.")
    actions.append("Click the original listing for every high-value or unusual figure. Confirm the price, status, fees, title and exact location with the advertiser.")
    actions.append("Use the Area Scorecard only where sample counts are adequate; commission additional checks where the sample is thin.")
    ai_actions=ai_commentary.get("recommended_actions") or []
    html_parts.extend(f"<li style=\"margin:7px 0\">{esc(x)}</li>" for x in (ai_actions[:5] or actions[:5]))
    html_parts.append("</ol>")
    if doc_url or sheet_url:
        html_parts.append("<h2 style=\"font-size:19px\">Full report and evidence</h2><p>The email is the executive briefing; use these links for the analyst report and property-level source trail.</p>")
        html_parts.append(link(doc_url,"Open full weekly report"))
        html_parts.append(link(sheet_url,"Open property data sheet"))
    html_parts.append("<p style=\"font-size:12px;color:#667085;border-top:1px solid #d9e1ea;padding-top:12px\">All figures are in Nigerian naira (₦). Amounts such as ₦450 million are rounded for readability; detailed sheets retain numeric values with comma separators. Prices are online asking prices, not confirmed transaction prices. Week-on-week percentages describe this tracked sample only. Missing URLs in the next scan do not, by themselves, prove a property has sold or been removed. Indicative gross yield excludes vacancy, service charges, repairs, taxes, financing and other costs.</p></main></body></html>")
    return "".join(html_parts)
def build_executive_html(listings_path,changes_path="none"):
    return build_executive_html_from_summary(_load_inputs(listings_path,changes_path)[-1])

def _xlsx_lib():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule
        return Workbook,Font,PatternFill,Alignment,Border,Side,get_column_letter,CellIsRule
    except ImportError as exc: raise RuntimeError("XLSX reporting requires openpyxl.") from exc

def build_xlsx_from_data(listings,changes,research,s,output=None):
    Workbook,Font,PatternFill,Alignment,Border,Side,get_column_letter,CellIsRule=_xlsx_lib()
    from openpyxl.chart import BarChart, Reference
    output=output or f"market_intelligence_{TODAY}.xlsx"
    wb=Workbook()
    navy="0D2340"; light="E8EEF5"; pale="F6F8FB"; green="E7F6EC"; red="FDECEC"; grey="667085"
    thin=Side(style="thin",color="D9E1EA")
    def title(sheet,text):
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=3)
        c=sheet.cell(1,1,text); c.font=Font(size=16,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=navy); c.alignment=Alignment(vertical="center")
        sheet.row_dimensions[1].height=30
    def headers(sheet,row=3):
        for c in sheet[row]:
            c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=navy); c.alignment=Alignment(wrap_text=True,vertical="center")
            c.border=Border(bottom=thin)
        sheet.freeze_panes=f"A{row+1}"
    def polish(sheet,start=1):
        for col in sheet.columns:
            letter=get_column_letter(col[0].column)
            maxlen=max((len(str(c.value or "")) for c in col),default=0)
            sheet.column_dimensions[letter].width=min(max(maxlen+2,12),38)
        for row in sheet.iter_rows(min_row=start):
            for c in row:
                c.alignment=Alignment(vertical="top",wrap_text=True)
        sheet.sheet_view.showGridLines=False
    ws=wb.active; ws.title="Executive Summary"
    title(ws,"Lagos & Abuja Property Market — Weekly Intelligence Brief")
    ws["A2"]=datetime.now(timezone.utc).strftime("%d %B %Y"); ws["A2"].font=Font(italic=True,color=grey)
    ws.append(["Metric","Current","Week-on-week"])
    rows=[("Listings tracked",s["listings_tracked"],"listings"),("Areas monitored",s["areas_monitored"],None),("Sources",s["sources"],None),
          ("New listings",s["changes"]["new_listings"],None),("Price reductions",s["changes"]["price_reductions"],None),("Price increases",s["changes"]["price_increases"],None),
          ("Median apartment sale",fmt_naira(s["median_apartment_sale"]),"sale"),("Median annual rent",fmt_naira(s["median_annual_rent"]),"rent"),("Median land ₦/sqm",fmt_naira(s["median_land_ppsqm"]),"land_ppsqm")]
    wow=s.get("week_on_week",{})
    for label,val,key in rows:
        w=_wow_text(wow.get("metrics",{}).get(key,{})) if key and wow.get("available") else "—"
        ws.append([label,val,w])
    headers(ws,3)
    ws["A14"]="Decision Signals"; ws["A14"].font=Font(size=13,bold=True,color=navy)
    for label,obj in [("Most expensive",s["signals"].get("most_expensive")),("Rental / investment screen",s["signals"].get("investment")),
                      ("Best buyer value screen",s["signals"].get("buyer_value")),("Lowest observed land ₦/sqm",s["signals"].get("land_opportunity"))]:
        ws.append([label,obj["market_node"] if obj else "N/A",""])
    ws.append(["Watch closely",", ".join(x["market_node"] for x in s["signals"].get("monitoring",[])) or "N/A",""])
    ws.append([]); ws.append(["Market interpretation"]); ws["A20"].font=Font(size=13,bold=True,color=navy)
    for note in s.get("analysis_notes",[]): ws.append([note])
    polish(ws)
    # Scorecard
    ws=wb.create_sheet("Area Scorecard")
    has_full_six_month=any(t.get("full_six_month") for t in s.get("trends",{}).values())
    trend_prefix="6-mo" if has_full_six_month else "Historical"
    ws.append(["Area","Listings","3BR Sale","3BR Rent","Yield Proxy","Sale ₦/sqm","Rent ₦/sqm","Land ₦/sqm",f"{trend_prefix} Sale",f"{trend_prefix} Rent",f"{trend_prefix} Land ₦/sqm","Sale Data","Rent Data","Land Data","3BR Sale n","3BR Rent n"])
    for x in s["scorecard"]:
        tr=s.get("trends",{}).get(x["market_node"],{})
        ws.append([x["market_node"],x["listing_count"],x["sale_3br_median"],x["rent_3br_median"],x["gross_rent_yield_proxy_pct"],x["sale_ppsqm_median"],x["rent_ppsqm_median"],x["land_ppsqm_median"],tr.get("sale_pct"),tr.get("rent_pct"),tr.get("land_ppsqm_pct"),x["sale_count"],x["rent_count"],x["land_count"],x.get("sale_3br_count",0),x.get("rent_3br_count",0)])
    headers(ws,1); polish(ws)

    # Visual trend chart: only chart observed historical percentages; never
    # synthesize missing periods or fill absent values with zero.
    trend_columns = (9, 10, 11)
    has_numeric_trend = any(
        isinstance(ws.cell(row=r, column=col).value, (int, float))
        for r in range(2, ws.max_row + 1)
        for col in trend_columns
    )
    if has_numeric_trend and ws.max_row > 1:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Observed Asking-Price Change by Area (%)"
        chart.y_axis.title = "Area"
        chart.x_axis.title = "Change (%)"
        chart.height = 8
        chart.width = 17
        data = Reference(ws, min_col=9, max_col=11, min_row=1, max_row=ws.max_row)
        categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.legend.position = "b"
        ws.add_chart(chart, "A12")

    # Current market snapshots are charted separately because sale prices,
    # annual rents, and land unit prices have very different scales.
    def add_area_chart(column, title_text, anchor):
        if not any(
            isinstance(ws.cell(row=r, column=column).value, (int, float))
            and ws.cell(row=r, column=column).value > 0
            for r in range(2, ws.max_row + 1)
        ):
            return
        current_chart = BarChart()
        current_chart.type = "bar"
        current_chart.style = 10
        current_chart.title = title_text
        current_chart.y_axis.title = "Area"
        current_chart.x_axis.title = "Naira (NGN)"
        current_chart.height = 7
        current_chart.width = 15
        current_data = Reference(
            ws, min_col=column, max_col=column, min_row=1, max_row=ws.max_row
        )
        current_categories = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        current_chart.add_data(current_data, titles_from_data=True)
        current_chart.set_categories(current_categories)
        current_chart.legend = None
        ws.add_chart(current_chart, anchor)

    add_area_chart(3, "Three-Bedroom Apartment Sale Asking Price by Area", "I12")
    add_area_chart(4, "Three-Bedroom Annual Rental Asking Price by Area", "A28")
    add_area_chart(8, "Land Asking Price per Square Metre by Area", "I28")

    # Data tabs
    def add_data(name,data):
        sh=wb.create_sheet(name)
        if data:
            sh.append(list(data[0].keys()))
            for r in data: sh.append([r.get(k) for k in data[0].keys()])
            headers(sh,1)
        else: sh.append(["No data returned"])
        polish(sh)
    add_data("Current Listings",listings); add_data("Weekly Changes",changes); add_data("Estate Intel Public",research)
    add_data("Additional Sources",SUPPLEMENTAL_SOURCE_ASSESSMENTS)
    source_ws=wb["Additional Sources"]
    for row in range(2, source_ws.max_row + 1):
        link_cell=source_ws.cell(row=row, column=2)
        if link_cell.value and str(link_cell.value).startswith(("https://","http://")):
            link_cell.hyperlink=str(link_cell.value)
            link_cell.style="Hyperlink"

    # A separate health tab makes source coverage and failures visible instead
    # of mixing scraper errors with market-price data.
    ws=wb.create_sheet("Source Health")
    ws.append(["Source","Output rows","Rows with public asking price","Rows with size/units","Failures","Status / interpretation"])
    try:
        run_report=json.loads(Path("scrape_run_report.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        run_report={}
    failures=run_report.get("failures",[]) if isinstance(run_report,dict) else []
    for source in ("Nigeria Property Centre","PropertyPro.ng","Estate Intel"):
        if source=="Estate Intel":
            source_rows=[r for r in research if (r.get("url") or "").startswith(("http://","https://"))]
            price_rows=sum(1 for r in source_rows if r.get("public_sale_price_ngn") not in (None,"") and str(r.get("public_sale_price_ngn")).strip())
            size_rows=sum(1 for r in source_rows if r.get("land_area_sqm") not in (None,"") or r.get("size_units") not in (None,""))
            note=("Public research only; no comparable listing prices inferred. Check linked pages and scrape failures."
                  if source_rows else "No public research records returned; review source-level failures.")
        else:
            source_rows=[r for r in listings if r.get("source")==source]
            price_rows=sum(1 for r in source_rows if r.get("asking_price_ngn") not in (None,"") and str(r.get("asking_price_ngn")).strip())
            size_rows=sum(1 for r in source_rows if r.get("size_sqm") not in (None,"") and str(r.get("size_sqm")).strip())
            note="Comparable listing rows; asking prices, not confirmed transactions."
        source_failures=[f for f in failures if f.get("source")==source]
        note += (f" {len(source_failures)} source/category failure(s) recorded." if source_failures else " No source/category failures recorded.")
        ws.append([source,len(source_rows),price_rows,size_rows,len(source_failures),note])
        for failure in source_failures:
            ws.append([source,0,0,0,1,f"{failure.get('node','Unknown node')} / {failure.get('category','research')}: {failure.get('error','Unknown error')}"])
    headers(ws,1); polish(ws)

    ws=wb.create_sheet("Methodology")
    for row in [["Purpose","Decision-ready property market intelligence"],["Comparable listing sources","Nigeria Property Centre; PropertyPro Nigeria"],["Research sources","Estate Intel public pages; Knight Frank Nigeria H1 2026 research reviewed as context; C&W Nigeria/Lagos-specific feed not yet verified"],
                ["Freshness","Current listing data targeted to 7 days; weekly refresh"],["Reporting & distribution","Scheduled Mondays at 08:00 Lagos time (07:00 UTC); email is attempted only after quality gates and publication, with provider receipt checked but inbox placement not guaranteed"],["Estate Intel commercial access","Not verifiable from this repository; confirm current agreement/subscription directly with account owner/vendor. Public pages only are used by this pipeline."],["Coverage","9 nodes: Banana Island, Old Ikoyi, Lekki Phase 1, Victoria Island, Eko Atlantic, Ikeja GRA, Asokoro, Maitama, Wuse"],
                ["Assets","Apartment sales/rentals 1–5BR; residential land; Estate Intel public research"],["Pricing","Asking prices; not confirmed closed transactions"],
                ["Estate Intel","Public research only; premium/login-gated values excluded"],["Transport","NPC direct; PropertyPro and Estate Intel via Zyte browser HTML"],
                ["History","Available-period directional comparison; labelled six-month only after at least 180 days of coverage"],["Decision use","Executive Summary for decisions; Area Scorecard for comparison; Current Listings for audit"]]: ws.append(row)
    headers(ws,1); polish(ws)
    # Column-aware scorecard formats. Trend values are percentage points, not
    # Excel ratios, so append a literal percent sign rather than multiplying by 100.
    score_ws = wb["Area Scorecard"]
    for col in ("C", "D", "F", "G", "H"):
        for cell in score_ws[col][1:]:
            cell.number_format = '#,##0'
    for col in ("E", "I", "J", "K"):
        for cell in score_ws[col][1:]:
            cell.number_format = '+0.0"%";-0.0"%";0.0"%"'
    for col in ("B", "L", "M", "N"):
        for cell in score_ws[col][1:]:
            cell.number_format = '#,##0'
    for sh in wb.worksheets:
        sh.auto_filter.ref=sh.dimensions if sh.max_row>1 else None
        if sh.title == "Area Scorecard":
            continue
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value,(int,float)) and c.column>1:
                    c.number_format='#,##0'
    wb.save(output); return output

def build_xlsx(listings_path,changes_path="none",research_path=None,output=None):
    return build_xlsx_from_data(*_load_inputs(listings_path,changes_path,research_path),output)

def build_docx_from_data(listings,changes,research,s,output=None):
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches,Pt
    except ImportError as exc: raise RuntimeError("DOCX reporting requires python-docx.") from exc
    output=output or f"Lagos_Abuja_Property_Market_Intelligence_{TODAY}.docx"
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.6); sec.left_margin=sec.right_margin=Inches(.7)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("LAGOS & ABUJA PROPERTY MARKET\nWEEKLY INTELLIGENCE REPORT"); r.bold=True; r.font.size=Pt(20)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(datetime.now(timezone.utc).strftime("%d %B %Y")).italic=True
    doc.add_heading("1. Executive Decision Brief",1)
    t=doc.add_table(rows=1,cols=3); t.style="Light Shading Accent 1"
    for i,h in enumerate(["Metric","Current","WoW"]): t.rows[0].cells[i].text=h
    wow=s.get("week_on_week",{})
    for label,val,key in [("Listings tracked",s["listings_tracked"],"listings"),("Areas monitored",s["areas_monitored"],None),("Sources",s["sources"],None),
                          ("New listings",s["changes"]["new_listings"],None),("Price reductions",s["changes"]["price_reductions"],None),("Price increases",s["changes"]["price_increases"],None),
                          ("Median apartment sale",fmt_naira(s["median_apartment_sale"]),"sale"),("Median annual rent",fmt_naira(s["median_annual_rent"]),"rent"),("Median land ₦/sqm",fmt_naira(s["median_land_ppsqm"]),"land_ppsqm")]:
        c=t.add_row().cells; c[0].text=label; c[1].text=str(val); c[2].text=_wow_text(wow.get("metrics",{}).get(key,{})) if key and wow.get("available") else "—"
    doc.add_heading("2. What Happened This Week",1)
    for n in s.get("analysis_notes",[]): doc.add_paragraph(n)
    doc.add_heading("3. Where the Opportunity Is",1)
    sig=s["signals"]
    for label,obj in [("Most expensive",sig.get("most_expensive")),("Strongest rental pricing / yield proxy",sig.get("strongest_rental")),("Best relative value",sig.get("best_relative_value")),("Lowest observed land ₦/sqm",sig.get("land_opportunity"))]:
        doc.add_paragraph(f"{label}: {obj['market_node'] if obj else 'N/A'}",style="List Bullet")
    doc.add_paragraph("Watch closely: "+(", ".join(x["market_node"] for x in sig.get("monitoring",[])) or "N/A"))
    doc.add_heading("4. Sales, Rentals & Land",1)
    for x in [z for z in s["scorecard"] if z["listing_count"]]:
        doc.add_paragraph(f"{x['market_node']}: {x['listing_count']} comparables; 3BR sale {fmt_naira(x['sale_3br_median'])}; 3BR rent {fmt_naira(x['rent_3br_median'])}; sale ₦/sqm {fmt_naira(x['sale_ppsqm_median'])}; land ₦/sqm {fmt_naira(x['land_ppsqm_median'])}.")
    full_six_month=any(tr.get("full_six_month") for tr in s.get("trends",{}).values())
    doc.add_heading("5. Six-Month Trend View" if full_six_month else "5. Historical Trend View",1)
    trend_rows=[(n,tr) for n,tr in s.get("trends",{}).items() if any(tr.get(k) is not None for k in ("sale_pct","rent_pct","land_ppsqm_pct"))]
    if trend_rows:
        for n,tr in trend_rows:
            vals=[]
            for key,label in (("sale_pct","sale"),("rent_pct","rent"),("land_ppsqm_pct","land ₦/sqm")):
                if tr.get(key) is not None: vals.append(f"{label} {tr[key]:+.1f}%")
            doc.add_paragraph(f"{n}: "+", ".join(vals),style="List Bullet")
    else: doc.add_paragraph("Trend baseline is being established; future weekly runs will extend the historical archive.")
    if not full_six_month and s.get("trends"):
        days=max((tr.get("period_days",0) for tr in s["trends"].values()), default=0)
        doc.add_paragraph(f"Available archived snapshots cover approximately {days} days; changes shown are directional for that available period, not a complete six-month series.")
    doc.add_heading("6. Recommended Actions",1)
    doc.add_paragraph("Use relative pricing and weekly movement to prioritise negotiations and areas for deeper diligence. Do not treat a thinly sampled area as a definitive market index.")
    doc.add_paragraph("All prices are asking prices. Estate Intel data is limited to public research; no premium or authentication controls are bypassed.")
    doc.save(output); return output

def build_docx(listings_path,changes_path="none",research_path=None,output=None):
    return build_docx_from_data(*_load_inputs(listings_path,changes_path,research_path),output)

def build_bundle(listings_path,changes_path="none",research_path=None):
    listings,changes,research,s=_load_inputs(listings_path,changes_path,research_path)
    html=build_executive_html_from_summary(s); html_path=f"executive_brief_{TODAY}.html"; Path(html_path).write_text(html,encoding="utf-8")
    xlsx=build_xlsx_from_data(listings,changes,research,s); docx=build_docx_from_data(listings,changes,research,s)
    analysis={"generated_at":datetime.now(timezone.utc).isoformat(),"summary":s,"outputs":{"html":html_path,"xlsx":xlsx,"docx":docx}}
    Path("market_intelligence.json").write_text(json.dumps(analysis,indent=2,ensure_ascii=False),encoding="utf-8")
    Path("report_outputs.json").write_text(json.dumps(analysis["outputs"],indent=2),encoding="utf-8")
    return analysis

def main():
    if len(sys.argv)<2: raise SystemExit("Usage: python report_builder.py <listings.csv> [changes.csv] [research.csv]")
    result=build_bundle(sys.argv[1],sys.argv[2] if len(sys.argv)>2 else "none",sys.argv[3] if len(sys.argv)>3 else "none")
    for k,v in result["outputs"].items(): print(f"Created {k}: {v}")

if __name__=="__main__": main()
